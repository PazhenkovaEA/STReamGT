"""Build a plate-template .xlsx for a kit: a pipeline-readable list + a linked plate grid.

Column A/B/C = TPositionId / SPositionBC / control_type (what make_ngsfilter reads, plain values;
control wells pre-filled). To the right, an 8x12 plate grid whose cells are formulas mirroring the
SPositionBC column, so filling the list updates the plate view. The file uploads back through the
job-submission "Upload Excel" path unchanged.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

LETTERS = "ABCDEFGH"
NUMS = list(range(1, 13))
GRID_COL0 = 5   # column E holds the plate row labels; F..Q hold columns 1..12

KIND_FILL = {
    "positive": "16A34A", "sequencing": "DC2626", "pcr": "F59E0B",
    "extraction": "7C3AED", "negative": "DC2626",
}


def _list_row(letter_idx: int, num: int) -> int:
    return 2 + letter_idx * 12 + (num - 1)   # rows 2..97, order A1..A12,B1..B12,...


def _plate_xlsx(named: dict[str, tuple[str | None, str]], note: str | None = None) -> bytes:
    """Build the plate .xlsx. `named`: well "A1" -> (sample_name, control_type). A control_type in
    KIND_FILL colours the well; empty control_type = an ordinary sample."""
    wb = Workbook()
    ws = wb.active
    ws.title = "plate"

    ws["A1"], ws["B1"], ws["C1"] = "Position", "Sample Name", "Control type"
    for cell in ("A1", "B1", "C1"):
        ws[cell].font = Font(bold=True)

    # list rows for all 96 wells; filled wells get name/control_type (+ colour for controls)
    for li, letter in enumerate(LETTERS):
        for num in NUMS:
            well = f"{letter}{num}"
            r = _list_row(li, num)
            ws.cell(row=r, column=1, value=well)
            entry = named.get(well)
            if entry:
                name, ct = entry
                nc = ws.cell(row=r, column=2, value=name)
                if ct:
                    ws.cell(row=r, column=3, value=ct)
                    nc.fill = PatternFill("solid", fgColor=KIND_FILL.get(ct, "DC2626"))
                    nc.font = Font(color="FFFFFF", bold=True)

    # plate grid header (well column numbers)
    for num in NUMS:
        hc = ws.cell(row=1, column=GRID_COL0 + num, value=num)
        hc.font = Font(bold=True)
        hc.alignment = Alignment(horizontal="center")

    # plate grid: each cell mirrors the list's SPositionBC via a formula
    for li, letter in enumerate(LETTERS):
        grow = 2 + li
        ws.cell(row=grow, column=GRID_COL0, value=letter).font = Font(bold=True)
        for num in NUMS:
            cell = ws.cell(row=grow, column=GRID_COL0 + num, value=f"=B{_list_row(li, num)}")
            cell.alignment = Alignment(horizontal="center")
            entry = named.get(f"{letter}{num}")
            if entry and entry[1]:
                cell.fill = PatternFill("solid", fgColor=KIND_FILL.get(entry[1], "DC2626"))
                cell.font = Font(color="FFFFFF", bold=True)

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 13
    for num in NUMS:
        ws.column_dimensions[get_column_letter(GRID_COL0 + num)].width = 13
    if note:
        ws.cell(row=11, column=GRID_COL0, value=note)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_control_template_xlsx(kit) -> bytes:
    """Blank plate template for a kit: control wells pre-filled, sample wells empty."""
    named = {c.position.upper(): (c.name, c.kind.value)
             for c in kit.controls if c.position}
    return _plate_xlsx(named, note=(
        "Fill sample names in the 'Sample Name' column (left); the plate grid mirrors them. "
        "Control rows are pre-filled — do not rename."))


def build_filled_plate_xlsx(rows) -> bytes:
    """A batch's submitted plate, filled: rows are {TPositionId, SPositionBC, control_type}."""
    named: dict[str, tuple[str | None, str]] = {}
    for r in rows:
        pos = str(r.get("TPositionId") or "").strip().upper()
        name = str(r.get("SPositionBC") or "").strip()
        if not pos or not name:
            continue
        named[pos] = (name, str(r.get("control_type") or "").strip().lower())
    return _plate_xlsx(named)
